from __future__ import annotations

import io
import json
import zipfile
import warnings
from datetime import datetime
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


warnings.filterwarnings("ignore")


# ---------------------------------------------------------
# Page configuration
# ---------------------------------------------------------

st.set_page_config(
    page_title="Reports & Export | BI Copilot",
    page_icon="📄",
    layout="wide",
)


st.markdown(
    """
    <style>
    .block-container {
        padding-top: 2rem;
        padding-bottom: 3rem;
        max-width: 1450px;
    }

    .page-header {
        padding: 1.8rem 2rem;
        border-radius: 20px;
        margin-bottom: 1.5rem;
        color: white;
        background: linear-gradient(
            120deg,
            #172554 0%,
            #1d4ed8 58%,
            #0ea5e9 100%
        );
        box-shadow: 0 16px 40px rgba(29, 78, 216, 0.18);
    }

    .page-header h1 {
        margin: 0;
        font-size: 2.1rem;
        font-weight: 750;
    }

    .page-header p {
        margin-top: 0.65rem;
        margin-bottom: 0;
        font-size: 1rem;
        opacity: 0.94;
        max-width: 950px;
    }

    .status-success {
        padding: 0.9rem 1rem;
        background: #ecfdf5;
        color: #065f46;
        border: 1px solid #a7f3d0;
        border-radius: 12px;
        margin-bottom: 1rem;
    }

    .status-info {
        padding: 0.9rem 1rem;
        background: #eff6ff;
        color: #1e3a8a;
        border: 1px solid #bfdbfe;
        border-radius: 12px;
        margin-bottom: 1rem;
    }

    div[data-testid="stMetric"] {
        border: 1px solid #dbe5f1;
        background: white;
        padding: 1rem;
        border-radius: 14px;
        box-shadow: 0 4px 14px rgba(15, 23, 42, 0.04);
    }
    </style>
    """,
    unsafe_allow_html=True,
)


st.markdown(
    """
    <div class="page-header">
        <h1>Professional Reports and Export</h1>
        <p>
            Consolidate data quality, statistical analysis, supervised
            learning, explainability, clustering and recommendations into
            downloadable PDF and PowerPoint reports.
        </p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ---------------------------------------------------------
# Data access
# ---------------------------------------------------------

def get_active_dataframe() -> pd.DataFrame | None:
    clean_df = st.session_state.get("clean_df")
    raw_df = st.session_state.get("raw_df")

    if isinstance(clean_df, pd.DataFrame) and not clean_df.empty:
        return clean_df.copy()

    if isinstance(raw_df, pd.DataFrame) and not raw_df.empty:
        return raw_df.copy()

    return None


df = get_active_dataframe()

if df is None:
    st.warning(
        "No active dataset was found. Return to the Home page and upload "
        "a dataset or select the demonstration dataset."
    )
    st.stop()


dataset_name = st.session_state.get(
    "dataset_name",
    "Current dataset",
)


st.markdown(
    f"""
    <div class="status-success">
        <b>Report dataset:</b> {dataset_name}
        &nbsp; | &nbsp;
        <b>{df.shape[0]:,}</b> rows
        &nbsp; | &nbsp;
        <b>{df.shape[1]:,}</b> columns
    </div>
    """,
    unsafe_allow_html=True,
)


# ---------------------------------------------------------
# Analysis summary
# ---------------------------------------------------------

def safe_float(
    value: Any,
    default: float | None = None,
) -> float | None:
    try:
        numeric_value = float(value)

        if np.isnan(numeric_value):
            return default

        return numeric_value
    except Exception:
        return default


def selected_model_information() -> dict[str, Any]:
    model_name = st.session_state.get(
        "active_model_name"
    )
    problem_type = st.session_state.get(
        "problem_type"
    )
    target_column = st.session_state.get(
        "target_column"
    )
    results = st.session_state.get(
        "model_results"
    )

    information = {
        "model_name": model_name,
        "problem_type": problem_type,
        "target_column": target_column,
        "metrics": {},
    }

    if (
        isinstance(results, pd.DataFrame)
        and not results.empty
        and model_name
    ):
        selected_rows = results.loc[
            results["model"] == model_name
        ]

        if not selected_rows.empty:
            row = selected_rows.iloc[0]

            if problem_type == "classification":
                information["metrics"] = {
                    "Accuracy": safe_float(
                        row.get("test_accuracy")
                    ),
                    "Balanced accuracy": safe_float(
                        row.get(
                            "test_balanced_accuracy"
                        )
                    ),
                    "Macro F1": safe_float(
                        row.get("test_f1_macro")
                    ),
                    "Weighted F1": safe_float(
                        row.get("test_f1_weighted")
                    ),
                    "ROC-AUC": safe_float(
                        row.get("test_roc_auc")
                    ),
                    "Log loss": safe_float(
                        row.get("test_log_loss")
                    ),
                    "CV macro F1": safe_float(
                        row.get("cv_f1_macro_mean")
                    ),
                }
            else:
                information["metrics"] = {
                    "MAE": safe_float(
                        row.get("test_mae")
                    ),
                    "RMSE": safe_float(
                        row.get("test_rmse")
                    ),
                    "R-squared": safe_float(
                        row.get("test_r2")
                    ),
                    "CV R-squared": safe_float(
                        row.get("cv_r2_mean")
                    ),
                }

    return information


def top_feature_names() -> list[str]:
    grouped_shap = st.session_state.get(
        "grouped_shap_df"
    )

    permutation_df = st.session_state.get(
        "permutation_importance_df"
    )

    grouped_builtin = st.session_state.get(
        "grouped_builtin_importance"
    )

    if (
        isinstance(grouped_shap, pd.DataFrame)
        and not grouped_shap.empty
    ):
        return (
            grouped_shap["original_feature"]
            .head(5)
            .astype(str)
            .tolist()
        )

    if (
        isinstance(permutation_df, pd.DataFrame)
        and not permutation_df.empty
    ):
        return (
            permutation_df["feature"]
            .head(5)
            .astype(str)
            .tolist()
        )

    if (
        isinstance(grouped_builtin, pd.DataFrame)
        and not grouped_builtin.empty
    ):
        return (
            grouped_builtin["original_feature"]
            .head(5)
            .astype(str)
            .tolist()
        )

    return []


def clustering_information() -> dict[str, Any]:
    cluster_count = st.session_state.get(
        "selected_cluster_count"
    )

    evaluation = st.session_state.get(
        "kmeans_evaluation"
    )

    anomaly_data = st.session_state.get(
        "anomaly_data"
    )

    clustered_data = st.session_state.get(
        "clustered_data"
    )

    silhouette = None

    if (
        cluster_count is not None
        and isinstance(evaluation, pd.DataFrame)
        and not evaluation.empty
    ):
        selected_rows = evaluation.loc[
            evaluation["number_of_clusters"]
            == cluster_count
        ]

        if not selected_rows.empty:
            silhouette = safe_float(
                selected_rows.iloc[0].get(
                    "silhouette_score"
                )
            )

    anomaly_count = (
        len(anomaly_data)
        if isinstance(anomaly_data, pd.DataFrame)
        else None
    )

    anomaly_percentage = None

    if (
        anomaly_count is not None
        and isinstance(clustered_data, pd.DataFrame)
        and len(clustered_data) > 0
    ):
        anomaly_percentage = (
            anomaly_count
            / len(clustered_data)
            * 100
        )

    return {
        "cluster_count": cluster_count,
        "silhouette_score": silhouette,
        "anomaly_count": anomaly_count,
        "anomaly_percentage": anomaly_percentage,
    }


def create_business_insights(
    dataframe: pd.DataFrame,
) -> list[str]:
    insights = [
        (
            f"The dataset contains {dataframe.shape[0]:,} records "
            f"and {dataframe.shape[1]:,} columns."
        )
    ]

    missing_count = int(
        dataframe.isna().sum().sum()
    )

    duplicate_count = int(
        dataframe.duplicated().sum()
    )

    if missing_count == 0:
        insights.append(
            "The active dataset has no missing values."
        )
    else:
        insights.append(
            f"The active dataset contains {missing_count:,} missing values."
        )

    if duplicate_count == 0:
        insights.append(
            "No duplicate rows were detected."
        )
    else:
        insights.append(
            f"The dataset contains {duplicate_count:,} duplicate rows."
        )

    model_info = selected_model_information()

    if model_info["model_name"]:
        insights.append(
            f"The validation-selected model is "
            f"{model_info['model_name']}."
        )

        if (
            model_info["problem_type"]
            == "classification"
        ):
            accuracy = model_info["metrics"].get(
                "Accuracy"
            )
            macro_f1 = model_info["metrics"].get(
                "Macro F1"
            )
            roc_auc = model_info["metrics"].get(
                "ROC-AUC"
            )

            metric_parts = []

            if accuracy is not None:
                metric_parts.append(
                    f"{accuracy:.2%} accuracy"
                )

            if macro_f1 is not None:
                metric_parts.append(
                    f"{macro_f1:.4f} macro F1"
                )

            if roc_auc is not None:
                metric_parts.append(
                    f"{roc_auc:.4f} ROC-AUC"
                )

            if metric_parts:
                insights.append(
                    "The selected model achieved "
                    + ", ".join(metric_parts)
                    + " on the holdout set."
                )

    top_features = top_feature_names()

    if top_features:
        insights.append(
            "The most influential model features are "
            + ", ".join(top_features)
            + "."
        )

    cluster_info = clustering_information()

    if cluster_info["cluster_count"] is not None:
        cluster_text = (
            f"K-Means selected "
            f"{cluster_info['cluster_count']} clusters."
        )

        if (
            cluster_info["silhouette_score"]
            is not None
        ):
            cluster_text += (
                f" The silhouette score is "
                f"{cluster_info['silhouette_score']:.4f}."
            )

        insights.append(cluster_text)

    if cluster_info["anomaly_count"] is not None:
        anomaly_text = (
            f"Isolation Forest identified "
            f"{cluster_info['anomaly_count']:,} "
            "potential anomalies"
        )

        if (
            cluster_info["anomaly_percentage"]
            is not None
        ):
            anomaly_text += (
                f", representing "
                f"{cluster_info['anomaly_percentage']:.2f}% "
                "of analyzed records"
            )

        insights.append(anomaly_text + ".")

    return insights


def create_recommendations() -> list[str]:
    recommendations = []

    top_features = top_feature_names()

    if top_features:
        recommendations.append(
            "Prioritize review of "
            + ", ".join(top_features[:3])
            + " because these features have the strongest influence "
            "on the current model."
        )

    model_info = selected_model_information()

    if model_info["model_name"]:
        recommendations.append(
            "Use the complete trained pipeline for prediction so that "
            "new data receives the same preprocessing used during training."
        )

    cluster_info = clustering_information()

    if (
        cluster_info["silhouette_score"]
        is not None
        and cluster_info["silhouette_score"] < 0.25
    ):
        recommendations.append(
            "Treat the clusters as exploratory because their low silhouette "
            "score indicates overlapping groups."
        )

    if cluster_info["anomaly_count"]:
        recommendations.append(
            "Review potential anomalies before removal because unusual "
            "records may represent data errors, rare events or valuable "
            "business cases."
        )

    recommendations.extend(
        [
            "Collect additional real-world observations before production deployment.",
            "Monitor model performance and data distributions for changes over time.",
            "Retrain and revalidate the model when business processes or data patterns change.",
        ]
    )

    return recommendations


def create_limitations() -> list[str]:
    limitations = [
        "Statistical associations, feature importance and SHAP values do not prove causation.",
        "Holdout performance may change when the model receives data from a different population or time period.",
        "Small datasets produce wider uncertainty and less stable estimates.",
        "Clustering labels represent statistical similarity and require business interpretation.",
        "Potential anomalies should be manually reviewed before deletion.",
        "Predictions should support business judgment rather than replace it.",
    ]

    return limitations


model_info = selected_model_information()
cluster_info = clustering_information()
business_insights = create_business_insights(df)
recommendations = create_recommendations()
limitations = create_limitations()
top_features = top_feature_names()


# ---------------------------------------------------------
# Report preview
# ---------------------------------------------------------

st.subheader("1. Report contents")

content_one, content_two, content_three, content_four = st.columns(4)

content_one.metric(
    "Dataset records",
    f"{len(df):,}",
)
content_two.metric(
    "Dataset columns",
    f"{df.shape[1]:,}",
)
content_three.metric(
    "Selected model",
    model_info["model_name"] or "Not trained",
)
content_four.metric(
    "Customer segments",
    cluster_info["cluster_count"]
    if cluster_info["cluster_count"] is not None
    else "Not generated",
)


preview_tab_one, preview_tab_two, preview_tab_three = st.tabs(
    [
        "Automated insights",
        "Recommendations",
        "Limitations",
    ]
)


with preview_tab_one:
    for index, insight in enumerate(
        business_insights,
        start=1,
    ):
        st.write(f"{index}. {insight}")


with preview_tab_two:
    for recommendation in recommendations:
        st.write(f"- {recommendation}")


with preview_tab_three:
    for limitation in limitations:
        st.write(f"- {limitation}")


# ---------------------------------------------------------
# Chart creation
# ---------------------------------------------------------

def create_numeric_summary_chart(
    dataframe: pd.DataFrame,
) -> bytes | None:
    numeric_columns = dataframe.select_dtypes(
        include=np.number
    ).columns.tolist()

    if not numeric_columns:
        return None

    selected_columns = numeric_columns[:6]

    means = (
        dataframe[selected_columns]
        .mean()
        .sort_values()
    )

    figure, axis = plt.subplots(
        figsize=(9, 4.5)
    )

    means.plot(
        kind="barh",
        ax=axis,
        color="#2563eb",
    )

    axis.set_title(
        "Average values of selected numeric features"
    )
    axis.set_xlabel("Mean")
    axis.grid(
        axis="x",
        alpha=0.2,
    )

    figure.tight_layout()

    image_file = io.BytesIO()

    figure.savefig(
        image_file,
        format="png",
        dpi=160,
        bbox_inches="tight",
    )

    plt.close(figure)

    image_file.seek(0)
    return image_file.getvalue()


def create_model_chart() -> bytes | None:
    results = st.session_state.get(
        "model_results"
    )

    problem_type = st.session_state.get(
        "problem_type"
    )

    if (
        not isinstance(results, pd.DataFrame)
        or results.empty
    ):
        return None

    if (
        problem_type == "classification"
        and "cv_f1_macro_mean" in results.columns
    ):
        values = (
            results.set_index("model")[
                "cv_f1_macro_mean"
            ]
            .sort_values()
        )

        axis_label = "Cross-validation macro F1"
        title = "Supervised model comparison"

    elif (
        problem_type == "regression"
        and "cv_r2_mean" in results.columns
    ):
        values = (
            results.set_index("model")[
                "cv_r2_mean"
            ]
            .sort_values()
        )

        axis_label = "Cross-validation R-squared"
        title = "Supervised model comparison"

    else:
        return None

    figure, axis = plt.subplots(
        figsize=(9, 5)
    )

    values.plot(
        kind="barh",
        ax=axis,
        color="#2563eb",
    )

    axis.set_title(title)
    axis.set_xlabel(axis_label)
    axis.grid(
        axis="x",
        alpha=0.2,
    )

    figure.tight_layout()

    image_file = io.BytesIO()

    figure.savefig(
        image_file,
        format="png",
        dpi=160,
        bbox_inches="tight",
    )

    plt.close(figure)

    image_file.seek(0)
    return image_file.getvalue()


def create_cluster_chart() -> bytes | None:
    cluster_profiles = st.session_state.get(
        "cluster_profiles"
    )

    if (
        not isinstance(cluster_profiles, pd.DataFrame)
        or cluster_profiles.empty
        or "records" not in cluster_profiles.columns
    ):
        return None

    figure, axis = plt.subplots(
        figsize=(9, 4.5)
    )

    axis.bar(
        cluster_profiles[
            "kmeans_cluster"
        ].astype(str),
        cluster_profiles["records"],
        color="#0ea5e9",
    )

    axis.set_title("Records in each K-Means cluster")
    axis.set_xlabel("Cluster")
    axis.set_ylabel("Records")
    axis.grid(
        axis="y",
        alpha=0.2,
    )

    figure.tight_layout()

    image_file = io.BytesIO()

    figure.savefig(
        image_file,
        format="png",
        dpi=160,
        bbox_inches="tight",
    )

    plt.close(figure)

    image_file.seek(0)
    return image_file.getvalue()


# ---------------------------------------------------------
# PDF generation
# ---------------------------------------------------------

def add_pdf_page_number(
    canvas: Any,
    document: Any,
) -> None:
    canvas.saveState()

    canvas.setFont(
        "Helvetica",
        8,
    )

    canvas.setFillColor(
        colors.HexColor("#64748b")
    )

    canvas.drawCentredString(
        A4[0] / 2,
        0.4 * inch,
        f"Intelligent BI Copilot | Page {document.page}",
    )

    canvas.restoreState()


def create_pdf_report() -> bytes:
    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=48,
        title="Intelligent Business Intelligence Copilot Report",
        author="Intelligent BI Copilot",
    )

    stylesheet = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CustomTitle",
        parent=stylesheet["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=28,
        textColor=colors.HexColor("#172554"),
        alignment=TA_CENTER,
        spaceAfter=18,
    )

    subtitle_style = ParagraphStyle(
        "CustomSubtitle",
        parent=stylesheet["Normal"],
        fontSize=10,
        leading=15,
        textColor=colors.HexColor("#475569"),
        alignment=TA_CENTER,
        spaceAfter=20,
    )

    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=stylesheet["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#1d4ed8"),
        spaceBefore=10,
        spaceAfter=10,
    )

    body_style = ParagraphStyle(
        "CustomBody",
        parent=stylesheet["BodyText"],
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=7,
    )

    small_style = ParagraphStyle(
        "SmallText",
        parent=body_style,
        fontSize=8,
        leading=11,
    )

    story = []

    story.append(
        Spacer(
            1,
            0.35 * inch,
        )
    )

    story.append(
        Paragraph(
            "Intelligent Business Intelligence Copilot",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Automated Data Quality, Statistical Analysis, Machine Learning, "
            "Explainability, Clustering and Business Insights",
            subtitle_style,
        )
    )

    cover_table = Table(
        [
            ["Dataset", str(dataset_name)],
            ["Generated", datetime.now().strftime("%d %B %Y, %H:%M")],
            ["Rows", f"{df.shape[0]:,}"],
            ["Columns", f"{df.shape[1]:,}"],
            [
                "Selected model",
                str(
                    model_info["model_name"]
                    or "Not trained"
                ),
            ],
            [
                "Target",
                str(
                    model_info["target_column"]
                    or "Not selected"
                ),
            ],
        ],
        colWidths=[
            1.7 * inch,
            4.3 * inch,
        ],
    )

    cover_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#dbeafe"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#172554"),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (1, 0),
                    (1, -1),
                    "Helvetica",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    9,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor("#bfdbfe"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ]
        )
    )

    story.append(cover_table)
    story.append(PageBreak())

    story.append(
        Paragraph(
            "1. Executive Summary",
            heading_style,
        )
    )

    for insight in business_insights:
        story.append(
            Paragraph(
                f"• {insight}",
                body_style,
            )
        )

    story.append(
        Paragraph(
            "2. Dataset and Data Quality",
            heading_style,
        )
    )

    missing_total = int(
        df.isna().sum().sum()
    )

    duplicate_total = int(
        df.duplicated().sum()
    )

    numeric_count = len(
        df.select_dtypes(include=np.number).columns
    )

    categorical_count = (
        df.shape[1] - numeric_count
    )

    quality_table = Table(
        [
            ["Measure", "Value"],
            ["Rows", f"{df.shape[0]:,}"],
            ["Columns", f"{df.shape[1]:,}"],
            ["Numeric columns", f"{numeric_count:,}"],
            ["Categorical/other columns", f"{categorical_count:,}"],
            ["Missing values", f"{missing_total:,}"],
            ["Duplicate rows", f"{duplicate_total:,}"],
        ],
        colWidths=[
            3.2 * inch,
            2.8 * inch,
        ],
    )

    quality_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#1d4ed8"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    9,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(quality_table)
    story.append(Spacer(1, 12))

    dtype_rows = [["Column", "Data type", "Missing", "Unique"]]

    for column in df.columns:
        dtype_rows.append(
            [
                str(column),
                str(df[column].dtype),
                str(int(df[column].isna().sum())),
                str(int(df[column].nunique(dropna=True))),
            ]
        )

    dtype_table = Table(
        dtype_rows,
        repeatRows=1,
        colWidths=[
            2.3 * inch,
            1.4 * inch,
            1.1 * inch,
            1.1 * inch,
        ],
    )

    dtype_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#e2e8f0"),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
            ]
        )
    )

    story.append(dtype_table)
    story.append(PageBreak())

    story.append(
        Paragraph(
            "3. Supervised Machine Learning",
            heading_style,
        )
    )

    if model_info["model_name"]:
        story.append(
            Paragraph(
                f"The validation-selected model is "
                f"<b>{model_info['model_name']}</b>. "
                f"The problem type is "
                f"<b>{str(model_info['problem_type']).title()}</b>, "
                f"and the target is "
                f"<b>{model_info['target_column']}</b>.",
                body_style,
            )
        )

        metric_rows = [["Metric", "Value"]]

        for metric_name, metric_value in (
            model_info["metrics"].items()
        ):
            if metric_value is not None:
                metric_rows.append(
                    [
                        metric_name,
                        f"{metric_value:.4f}",
                    ]
                )

        metric_table = Table(
            metric_rows,
            colWidths=[
                3.2 * inch,
                2.8 * inch,
            ],
        )

        metric_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#1d4ed8"),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.4,
                        colors.HexColor("#cbd5e1"),
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        9,
                    ),
                ]
            )
        )

        story.append(metric_table)

    else:
        story.append(
            Paragraph(
                "No supervised model was trained during this session.",
                body_style,
            )
        )

    story.append(
        Paragraph(
            "4. Explainable AI",
            heading_style,
        )
    )

    if top_features:
        story.append(
            Paragraph(
                "The most influential features are: "
                + ", ".join(top_features)
                + ".",
                body_style,
            )
        )
    else:
        story.append(
            Paragraph(
                "Feature importance was not generated during this session.",
                body_style,
            )
        )

    story.append(
        Paragraph(
            "Feature importance and SHAP values explain model associations. "
            "They do not prove that changing a feature will cause the "
            "predicted outcome to change.",
            body_style,
        )
    )

    story.append(
        Paragraph(
            "5. Unsupervised Learning",
            heading_style,
        )
    )

    if cluster_info["cluster_count"] is not None:
        story.append(
            Paragraph(
                f"K-Means selected "
                f"<b>{cluster_info['cluster_count']}</b> clusters.",
                body_style,
            )
        )

        if (
            cluster_info["silhouette_score"]
            is not None
        ):
            story.append(
                Paragraph(
                    f"The silhouette score was "
                    f"<b>{cluster_info['silhouette_score']:.4f}</b>.",
                    body_style,
                )
            )

        if cluster_info["anomaly_count"] is not None:
            story.append(
                Paragraph(
                    f"Isolation Forest identified "
                    f"<b>{cluster_info['anomaly_count']:,}</b> "
                    "potential anomalies.",
                    body_style,
                )
            )
    else:
        story.append(
            Paragraph(
                "Unsupervised learning was not run during this session.",
                body_style,
            )
        )

    story.append(PageBreak())

    story.append(
        Paragraph(
            "6. Recommendations",
            heading_style,
        )
    )

    for recommendation in recommendations:
        story.append(
            Paragraph(
                f"• {recommendation}",
                body_style,
            )
        )

    story.append(
        Paragraph(
            "7. Limitations",
            heading_style,
        )
    )

    for limitation in limitations:
        story.append(
            Paragraph(
                f"• {limitation}",
                body_style,
            )
        )

    story.append(
        Spacer(
            1,
            15,
        )
    )

    story.append(
        Paragraph(
            "This report was generated automatically by the Intelligent "
            "Business Intelligence Copilot using local and open-source "
            "Python tools.",
            small_style,
        )
    )

    document.build(
        story,
        onFirstPage=add_pdf_page_number,
        onLaterPages=add_pdf_page_number,
    )

    output.seek(0)
    return output.getvalue()


# ---------------------------------------------------------
# PowerPoint generation
# ---------------------------------------------------------

def set_slide_title(
    slide: Any,
    title: str,
) -> None:
    title_shape = slide.shapes.title
    title_shape.text = title

    paragraph = title_shape.text_frame.paragraphs[0]
    paragraph.font.name = "Aptos Display"
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(
        23,
        37,
        84,
    )


def add_bullet_slide(
    presentation: Presentation,
    title: str,
    bullet_items: list[str],
) -> None:
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    set_slide_title(
        slide,
        title,
    )

    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()

    for item_index, item in enumerate(
        bullet_items
    ):
        paragraph = (
            text_frame.paragraphs[0]
            if item_index == 0
            else text_frame.add_paragraph()
        )

        paragraph.text = str(item)
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = RGBColor(
            30,
            41,
            59,
        )
        paragraph.space_after = Pt(8)


def add_chart_slide(
    presentation: Presentation,
    title: str,
    image_bytes: bytes,
) -> None:
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    set_slide_title(
        slide,
        title,
    )

    image_file = io.BytesIO(
        image_bytes
    )

    slide.shapes.add_picture(
        image_file,
        Inches(1.05),
        Inches(1.45),
        width=Inches(11.2),
        height=Inches(5.4),
    )


def create_powerpoint_report() -> bytes:
    presentation = Presentation()

    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    title_slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    title_slide.shapes.title.text = (
        "Intelligent Business Intelligence Copilot"
    )

    title_paragraph = (
        title_slide.shapes.title.text_frame.paragraphs[0]
    )

    title_paragraph.font.name = "Aptos Display"
    title_paragraph.font.size = Pt(30)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = RGBColor(
        23,
        37,
        84,
    )
    title_paragraph.alignment = PP_ALIGN.CENTER

    subtitle = title_slide.placeholders[1]
    subtitle.text = (
        f"Automated Analytical Report\n"
        f"Dataset: {dataset_name}\n"
        f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}"
    )

    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = RGBColor(
            71,
            85,
            105,
        )
        paragraph.alignment = PP_ALIGN.CENTER

    add_bullet_slide(
        presentation,
        "Executive Summary",
        business_insights,
    )

    quality_bullets = [
        f"Rows: {df.shape[0]:,}",
        f"Columns: {df.shape[1]:,}",
        (
            "Numeric columns: "
            f"{len(df.select_dtypes(include=np.number).columns):,}"
        ),
        (
            "Missing values: "
            f"{int(df.isna().sum().sum()):,}"
        ),
        (
            "Duplicate rows: "
            f"{int(df.duplicated().sum()):,}"
        ),
    ]

    add_bullet_slide(
        presentation,
        "Dataset and Data Quality",
        quality_bullets,
    )

    numeric_chart = create_numeric_summary_chart(
        df
    )

    if numeric_chart is not None:
        add_chart_slide(
            presentation,
            "Numeric Feature Overview",
            numeric_chart,
        )

    if model_info["model_name"]:
        model_bullets = [
            (
                f"Selected model: "
                f"{model_info['model_name']}"
            ),
            (
                f"Problem type: "
                f"{str(model_info['problem_type']).title()}"
            ),
            (
                f"Target feature: "
                f"{model_info['target_column']}"
            ),
        ]

        for metric_name, metric_value in (
            model_info["metrics"].items()
        ):
            if metric_value is not None:
                model_bullets.append(
                    f"{metric_name}: "
                    f"{metric_value:.4f}"
                )

        add_bullet_slide(
            presentation,
            "Supervised Machine Learning",
            model_bullets,
        )

        model_chart = create_model_chart()

        if model_chart is not None:
            add_chart_slide(
                presentation,
                "Model Comparison",
                model_chart,
            )

    if top_features:
        explainability_bullets = [
            "Most influential features: "
            + ", ".join(top_features),
            (
                "Feature importance explains predictive influence, "
                "not causation."
            ),
            (
                "Use the full preprocessing and model pipeline for "
                "new predictions."
            ),
        ]

        add_bullet_slide(
            presentation,
            "Explainable AI",
            explainability_bullets,
        )

    if cluster_info["cluster_count"] is not None:
        clustering_bullets = [
            (
                f"Selected K-Means clusters: "
                f"{cluster_info['cluster_count']}"
            ),
        ]

        if (
            cluster_info["silhouette_score"]
            is not None
        ):
            clustering_bullets.append(
                f"Silhouette score: "
                f"{cluster_info['silhouette_score']:.4f}"
            )

        if cluster_info["anomaly_count"] is not None:
            clustering_bullets.append(
                f"Potential anomalies: "
                f"{cluster_info['anomaly_count']:,}"
            )

        clustering_bullets.append(
            "Cluster names require business validation."
        )

        add_bullet_slide(
            presentation,
            "Unsupervised Learning",
            clustering_bullets,
        )

        cluster_chart = create_cluster_chart()

        if cluster_chart is not None:
            add_chart_slide(
                presentation,
                "Cluster Distribution",
                cluster_chart,
            )

    add_bullet_slide(
        presentation,
        "Recommendations",
        recommendations,
    )

    add_bullet_slide(
        presentation,
        "Limitations",
        limitations,
    )

    add_bullet_slide(
        presentation,
        "Conclusion",
        [
            (
                "The BI Copilot completed an integrated workflow covering "
                "data quality, statistics, machine learning, explainability "
                "and reporting."
            ),
            (
                "The selected model and discovered segments should be "
                "validated with current real-world business data."
            ),
            (
                "Continue monitoring data quality, model performance and "
                "distribution changes after deployment."
            ),
        ],
    )

    output = io.BytesIO()
    presentation.save(output)
    output.seek(0)

    return output.getvalue()


# ---------------------------------------------------------
# JSON and ZIP package
# ---------------------------------------------------------

def create_analysis_summary() -> dict[str, Any]:
    return {
        "generated_at": datetime.now().isoformat(),
        "dataset": {
            "name": dataset_name,
            "rows": int(df.shape[0]),
            "columns": int(df.shape[1]),
            "column_names": list(
                map(str, df.columns)
            ),
            "missing_values": int(
                df.isna().sum().sum()
            ),
            "duplicate_rows": int(
                df.duplicated().sum()
            ),
            "numeric_columns": (
                df.select_dtypes(
                    include=np.number
                )
                .columns
                .astype(str)
                .tolist()
            ),
        },
        "supervised_model": model_info,
        "important_features": top_features,
        "unsupervised_analysis": cluster_info,
        "business_insights": business_insights,
        "recommendations": recommendations,
        "limitations": limitations,
    }


def create_zip_package(
    pdf_bytes: bytes,
    powerpoint_bytes: bytes,
    summary_json: bytes,
) -> bytes:
    output = io.BytesIO()

    with zipfile.ZipFile(
        output,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        archive.writestr(
            "Intelligent_BI_Copilot_Report.pdf",
            pdf_bytes,
        )

        archive.writestr(
            "Intelligent_BI_Copilot_Presentation.pptx",
            powerpoint_bytes,
        )

        archive.writestr(
            "analysis_summary.json",
            summary_json,
        )

        archive.writestr(
            "active_dataset.csv",
            df.to_csv(
                index=False
            ).encode("utf-8"),
        )

        model_results = st.session_state.get(
            "model_results"
        )

        if isinstance(
            model_results,
            pd.DataFrame,
        ):
            archive.writestr(
                "supervised_model_results.csv",
                model_results.to_csv(
                    index=False
                ).encode("utf-8"),
            )

        predictions = st.session_state.get(
            "model_predictions"
        )

        if isinstance(
            predictions,
            pd.DataFrame,
        ):
            archive.writestr(
                "supervised_test_predictions.csv",
                predictions.to_csv(
                    index=False
                ).encode("utf-8"),
            )

        cluster_profiles = st.session_state.get(
            "cluster_profiles"
        )

        if isinstance(
            cluster_profiles,
            pd.DataFrame,
        ):
            archive.writestr(
                "cluster_profiles.csv",
                cluster_profiles.to_csv(
                    index=False
                ).encode("utf-8"),
            )

        anomaly_data = st.session_state.get(
            "anomaly_data"
        )

        if isinstance(
            anomaly_data,
            pd.DataFrame,
        ):
            archive.writestr(
                "potential_anomalies.csv",
                anomaly_data.to_csv(
                    index=False
                ).encode("utf-8"),
            )

        shap_importance = st.session_state.get(
            "shap_importance_df"
        )

        if isinstance(
            shap_importance,
            pd.DataFrame,
        ):
            archive.writestr(
                "global_shap_importance.csv",
                shap_importance.to_csv(
                    index=False
                ).encode("utf-8"),
            )

        permutation_data = st.session_state.get(
            "permutation_importance_df"
        )

        if isinstance(
            permutation_data,
            pd.DataFrame,
        ):
            archive.writestr(
                "permutation_importance.csv",
                permutation_data.to_csv(
                    index=False
                ).encode("utf-8"),
            )

    output.seek(0)
    return output.getvalue()


# ---------------------------------------------------------
# Generate reports
# ---------------------------------------------------------

st.subheader("2. Generate professional reports")

st.markdown(
    """
    <div class="status-info">
        Reports are generated locally. The active dataset is not sent to
        an external paid reporting service.
    </div>
    """,
    unsafe_allow_html=True,
)


if st.button(
    "Generate PDF and PowerPoint reports",
    type="primary",
    width="stretch",
):
    with st.spinner(
        "Generating PDF, PowerPoint and analysis package..."
    ):
        try:
            pdf_report = create_pdf_report()
            powerpoint_report = (
                create_powerpoint_report()
            )

            analysis_summary = (
                create_analysis_summary()
            )

            summary_json = json.dumps(
                analysis_summary,
                indent=2,
                default=str,
            ).encode("utf-8")

            zip_package = create_zip_package(
                pdf_report,
                powerpoint_report,
                summary_json,
            )

            st.session_state.generated_pdf_report = (
                pdf_report
            )

            st.session_state.generated_powerpoint_report = (
                powerpoint_report
            )

            st.session_state.generated_summary_json = (
                summary_json
            )

            st.session_state.generated_analysis_package = (
                zip_package
            )

            st.success(
                "PDF, PowerPoint and analysis package were generated "
                "successfully."
            )

        except Exception as error:
            st.error(
                f"Report generation failed: {error}"
            )


pdf_report = st.session_state.get(
    "generated_pdf_report"
)

powerpoint_report = st.session_state.get(
    "generated_powerpoint_report"
)

summary_json = st.session_state.get(
    "generated_summary_json"
)

analysis_package = st.session_state.get(
    "generated_analysis_package"
)


if pdf_report and powerpoint_report:
    st.subheader("3. Download generated files")

    download_one, download_two = st.columns(2)

    with download_one:
        st.download_button(
            "Download PDF report",
            data=pdf_report,
            file_name=(
                "Intelligent_BI_Copilot_Report.pdf"
            ),
            mime="application/pdf",
            width="stretch",
        )

    with download_two:
        st.download_button(
            "Download PowerPoint presentation",
            data=powerpoint_report,
            file_name=(
                "Intelligent_BI_Copilot_Presentation.pptx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            width="stretch",
        )

    download_three, download_four = st.columns(2)

    with download_three:
        st.download_button(
            "Download analysis summary",
            data=summary_json,
            file_name="analysis_summary.json",
            mime="application/json",
            width="stretch",
        )

    with download_four:
        st.download_button(
            "Download complete analysis package",
            data=analysis_package,
            file_name=(
                "Intelligent_BI_Copilot_Analysis_Package.zip"
            ),
            mime="application/zip",
            width="stretch",
        )

    file_summary = pd.DataFrame(
        {
            "file": [
                "Intelligent_BI_Copilot_Report.pdf",
                (
                    "Intelligent_BI_Copilot_"
                    "Presentation.pptx"
                ),
                "analysis_summary.json",
                (
                    "Intelligent_BI_Copilot_"
                    "Analysis_Package.zip"
                ),
            ],
            "size_bytes": [
                len(pdf_report),
                len(powerpoint_report),
                len(summary_json),
                len(analysis_package),
            ],
        }
    )

    file_summary["size_kb"] = (
        file_summary["size_bytes"] / 1024
    ).round(2)

    st.dataframe(
        file_summary,
        width="stretch",
        hide_index=True,
    )

else:
    st.info(
        "Select **Generate PDF and PowerPoint reports** to create the "
        "downloadable deliverables."
    )