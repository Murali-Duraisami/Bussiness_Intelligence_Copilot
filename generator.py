from pathlib import Path
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from reportlab.lib import colors
from pptx import Presentation

def pdf_report(path,title,summary,insights):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); styles=getSampleStyleSheet(); story=[Paragraph(title,styles['Title']),Spacer(1,12)]
    story += [Paragraph(f'<b>{k}:</b> {v}',styles['BodyText']) for k,v in summary.items()]
    story += [Spacer(1,12),Paragraph('Key insights',styles['Heading2'])]
    story += [Paragraph('• '+str(i),styles['BodyText']) for i in insights]
    SimpleDocTemplate(str(path),pagesize=A4).build(story); return path

def ppt_report(path,title,summary,insights):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text=title; slide.placeholders[1].text='Automated Business Intelligence Report'
    for heading,items in [('Dataset summary',[f'{k}: {v}' for k,v in summary.items()]),('Key insights',insights)]:
        s=prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text=heading; s.placeholders[1].text='\n'.join('• '+str(x) for x in items)
    prs.save(path); return path
